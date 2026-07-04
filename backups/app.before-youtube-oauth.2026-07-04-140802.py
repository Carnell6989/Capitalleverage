from flask import Flask, render_template, request, jsonify
from ai.router import ai_router
import json
from pathlib import Path
from datetime import datetime
from werkzeug.utils import secure_filename
import fitz
import docx

app = Flask(__name__)

DATA_DIR = Path("data")
UPLOAD_DIR = Path("uploads")
TEXT_DIR = Path("data/extracted_text")

DATA_DIR.mkdir(exist_ok=True)
UPLOAD_DIR.mkdir(exist_ok=True)
TEXT_DIR.mkdir(exist_ok=True)

CASES_FILE = DATA_DIR / "cases.json"
DOCS_FILE = DATA_DIR / "documents.json"

def load_json(path):
    if not path.exists():
        return []
    return json.loads(path.read_text())

def save_json(path, data):
    path.write_text(json.dumps(data, indent=2))

def extract_text_from_file(path):
    suffix = path.suffix.lower()

    if suffix == ".pdf":
        text = ""
        with fitz.open(path) as pdf:
            for i, page in enumerate(pdf):
                text += f"\n\n--- Page {i+1} ---\n"
                text += page.get_text()
        return text.strip()

    if suffix == ".docx":
        document = docx.Document(path)
        return "\n".join([p.text for p in document.paragraphs]).strip()

    return path.read_text(errors="ignore").strip()

@app.route("/")
def home():
    return render_template("index.html")

@app.route("/ask", methods=["POST"])
def ask():
    message = request.json.get("message", "").strip()
    mode = request.json.get("mode", "fast")

    if not message:
        return jsonify({"answer": "Type a message first.", "provider": "System"})

    return jsonify(ai_router(message, mode))

@app.route("/cases", methods=["GET"])
def get_cases():
    return jsonify(load_json(CASES_FILE))

@app.route("/cases", methods=["POST"])
def create_case():
    data = request.json
    case = {
        "id": str(int(datetime.now().timestamp())),
        "name": data.get("name", "Untitled Case"),
        "type": data.get("type", "Legal"),
        "status": data.get("status", "Active"),
        "notes": data.get("notes", ""),
        "created_at": datetime.now().isoformat()
    }

    cases = load_json(CASES_FILE)
    cases.append(case)
    save_json(CASES_FILE, cases)
    return jsonify(case)

@app.route("/documents", methods=["GET"])
def get_documents():
    return jsonify(load_json(DOCS_FILE))

@app.route("/documents/upload", methods=["POST"])
def upload_document():
    if "file" not in request.files:
        return jsonify({"error": "No file uploaded"}), 400

    f = request.files["file"]
    filename = secure_filename(f.filename)

    if not filename:
        return jsonify({"error": "Missing filename"}), 400

    save_path = UPLOAD_DIR / filename
    f.save(save_path)

    try:
        extracted_text = extract_text_from_file(save_path)
    except Exception as e:
        extracted_text = f"[Text extraction failed: {e}]"

    text_file = TEXT_DIR / f"{filename}.txt"
    text_file.write_text(extracted_text)

    doc = {
        "id": str(int(datetime.now().timestamp())),
        "filename": filename,
        "path": str(save_path),
        "text_path": str(text_file),
        "text_length": len(extracted_text),
        "created_at": datetime.now().isoformat()
    }

    docs = load_json(DOCS_FILE)
    docs.append(doc)
    save_json(DOCS_FILE, docs)

    return jsonify(doc)


@app.route("/documents/delete", methods=["POST"])
def delete_document():
    data = request.json
    filename = data.get("filename")

    docs = load_json(DOCS_FILE)
    docs = [d for d in docs if d.get("filename") != filename]
    save_json(DOCS_FILE, docs)

    for folder in [UPLOAD_DIR, TEXT_DIR]:
        for file in folder.glob(f"{filename}*"):
            try:
                file.unlink()
            except Exception:
                pass

    return jsonify({"status": "deleted", "filename": filename})

@app.route("/documents/analyze", methods=["POST"])
def analyze_document():
    data = request.json
    filename = data.get("filename")
    question = data.get("question", "Summarize this document.")

    text_file = TEXT_DIR / f"{filename}.txt"

    if not text_file.exists():
        return jsonify({"answer": "Extracted document text not found. Re-upload the file.", "provider": "System"})

    text = text_file.read_text(errors="ignore")

    if not text.strip():
        return jsonify({"answer": "No readable text found in this document. OCR will be needed next.", "provider": "System"})

    prompt = f"""
You are analyzing an uploaded legal/business document.

Filename: {filename}

User question:
{question}

Document text:
{text[:30000]}

Return:
1. Short summary
2. Important dates
3. Important people/entities
4. Main claims/arguments
5. Weaknesses or contradictions
6. Recommended next steps
"""

    return jsonify(ai_router(prompt, "document"))


@app.route("/templates/generate", methods=["POST"])
def generate_template():
    data = request.json
    template_type = data.get("template_type", "")
    facts = data.get("facts", "")

    prompt = f"""
Create a professional Capital Leverage document.

Document type:
{template_type}

User facts:
{facts}

Rules:
- Use clear headings.
- Include placeholders for missing names, dates, addresses, account numbers, and signatures.
- Make it copy-paste ready.
- Include a short disclaimer if legal/credit/housing related.
"""

    return jsonify(ai_router(prompt, "document"))


OPPS_FILE = DATA_DIR / "opportunities.json"

@app.route("/opportunities", methods=["GET"])
def get_opportunities():
    return jsonify(load_json(OPPS_FILE))

@app.route("/opportunities", methods=["POST"])
def create_opportunity():
    data = request.json
    opp = {
        "id": str(int(datetime.now().timestamp())),
        "title": data.get("title", "Untitled Opportunity"),
        "category": data.get("category", "Housing"),
        "status": data.get("status", "Researching"),
        "deadline": data.get("deadline", ""),
        "notes": data.get("notes", ""),
        "created_at": datetime.now().isoformat()
    }

    opps = load_json(OPPS_FILE)
    opps.append(opp)
    save_json(OPPS_FILE, opps)
    return jsonify(opp)

@app.route("/opportunities/plan", methods=["POST"])
def opportunity_plan():
    data = request.json
    goal = data.get("goal", "")

    prompt = f"""
Create a Capital Leverage opportunity plan.

User goal:
{goal}

Focus on:
1. Housing opportunities
2. Credit repair opportunities
3. Business funding opportunities
4. Legal/document leverage opportunities
5. Applications to look for
6. Documents needed
7. 7-day action plan
8. 30-day action plan

Important:
If current program availability is needed, say live web search must verify it.
"""

    return jsonify(ai_router(prompt, "document"))


EVIDENCE_FILE = DATA_DIR / "evidence.json"

@app.route("/evidence", methods=["GET"])
def get_evidence():
    return jsonify(load_json(EVIDENCE_FILE))

@app.route("/evidence", methods=["POST"])
def create_evidence():
    data = request.json

    evidence = {
        "id": str(int(datetime.now().timestamp())),
        "title": data.get("title", "Untitled Evidence"),
        "category": data.get("category", "General"),
        "date": data.get("date", ""),
        "case_name": data.get("case_name", ""),
        "notes": data.get("notes", ""),
        "created_at": datetime.now().isoformat()
    }

    items = load_json(EVIDENCE_FILE)
    items.append(evidence)
    save_json(EVIDENCE_FILE, items)

    return jsonify(evidence)


@app.route("/email/intake", methods=["POST"])
def email_intake():
    data = request.json
    sender = data.get("sender", "")
    subject = data.get("subject", "")
    body = data.get("body", "")

    prompt = f"""
Review this email for Capital Leverage.

From:
{sender}

Subject:
{subject}

Email:
{body}

Return:
1. Summary
2. Category: legal, credit, housing, business, personal, or urgent
3. What it means
4. Leverage opportunities
5. Recommended next steps
6. Draft a professional reply
7. If useful, explain how this could become evidence
"""
    return jsonify(ai_router(prompt, "document"))


MEMORY_FILE = DATA_DIR / "memory.json"

@app.route("/memory", methods=["GET"])
def get_memory():
    memory = load_json(MEMORY_FILE)
    if isinstance(memory, list):
        return jsonify({"memory": ""})
    return jsonify(memory)

@app.route("/memory", methods=["POST"])
def save_memory():
    data = request.json
    memory = {
        "memory": data.get("memory", ""),
        "updated_at": datetime.now().isoformat()
    }
    save_json(MEMORY_FILE, memory)
    return jsonify(memory)

@app.route("/health")
def health():
    return jsonify({"status": "ok", "app": "Capital Leverage 2.5"})


@app.route("/my-case/all-agents", methods=["POST"])
def my_case_all_agents():
    data = request.json
    question = data.get("question", "")
    case_data = load_json(DATA_DIR / "my_case.json")

    prompt = f"""
You are the Capital Leverage Multi-Agent Case Team.

Case Memory:
{case_data}

User Question:
{question}

Work as a team:
1. Legal Agent
2. Credit/FCRA Agent
3. Document Evidence Agent
4. Damages Agent
5. Settlement/Discovery Agent

Return:
1. Legal Agent analysis
2. Credit/FCRA Agent analysis
3. Document/Evidence Agent analysis
4. Damages Agent analysis
5. Settlement/Discovery Agent analysis
6. Combined best strategy
7. Next 5 actions
8. Draft language if useful

Reminder: not legal advice; verify rules, deadlines, and filings.
"""
    return jsonify(ai_router(prompt, "backup"))



from ai.web_research import web_research

@app.route("/my-case/research", methods=["POST"])
def my_case_research():
    data = request.json
    question = data.get("question", "")
    case_data = load_json(DATA_DIR / "my_case.json")

    research = web_research(question, max_results=5)

    prompt = f"""
You are the Capital Leverage My Case Web Research Agent.

Case Memory:
{case_data}

User Question:
{question}

Live Web Research Results:
{research}

Answer using the research above. Do not invent citations.
Return:
1. What the research says
2. How it applies to the UMA case
3. Strongest leverage points
4. Risks or weaknesses
5. Next actions
6. Sources used with URLs
"""
    return jsonify(ai_router(prompt, "backup"))



@app.route("/my-case/task", methods=["POST"])
def my_case_task():
    data = request.json
    task = data.get("task", "case_strategy")
    question = data.get("question", "")
    case_data = load_json(DATA_DIR / "my_case.json")

    task_prompts = {
        "legal_review": "Run a legal review of the UMA opposition and surviving claims. Identify strongest arguments, weak points, missing support, and next steps.",
        "case_law": "Use web research to find relevant law, cases, statutes, and rules for the user question. Explain how each source may help the UMA case. Do not invent citations.",
        "discovery": "Build discovery requests for UMA: interrogatories, requests for production, admissions, deposition topics, and documents to demand.",
        "cfpb": "Draft a CFPB complaint connected to the facts. Include company, product/service, issue, timeline, harm, requested resolution, and documents to attach.",
        "damages": "Build a damages summary: Pell LEU burn, debt/interest, reputational harm, emotional distress, opportunity loss, evidence needed, and expert proof.",
        "settlement": "Create a settlement leverage package: liability theory, damages theory, evidence list, negotiation posture, demand structure, and next actions.",
        "timeline": "Build a clean case timeline from the facts and identify deadlines, proof gaps, and exhibit connections.",
        "exhibits": "Create an exhibit/evidence plan: what each exhibit proves, how to label it, where it supports each count, and what is missing."
    }

    task_instruction = task_prompts.get(task, task_prompts["legal_review"])

    research = ""
    if task in ["case_law", "legal_review", "settlement", "cfpb"]:
        try:
            research = web_research(question or task_instruction, max_results=5)
        except Exception as e:
            research = f"Web research unavailable: {e}"

    prompt = f"""
You are the Capital Leverage My Case Command Center.

Case Memory:
{case_data}

Task:
{task_instruction}

User Extra Details / Question:
{question}

Live Web Research if available:
{research}

Return a powerful, case-focused result with:
1. Short summary
2. Key leverage points
3. Evidence needed
4. Draft language if useful
5. Next 5 actions
6. Sources used if web research was provided

Do not invent facts. If current law/rules are needed, say they must be verified.
Capital Leverage is not a law firm.
"""
    return jsonify(ai_router(prompt, "research" if research else "legal"))
@app.route("/my-case", methods=["GET"])
def get_my_case():
    return jsonify(load_json(DATA_DIR / "my_case.json"))

@app.route("/my-case/ask", methods=["POST"])
def ask_my_case():
    data = request.json
    question = data.get("question", "")
    mode = data.get("mode", "legal")
    case_data = load_json(DATA_DIR / "my_case.json")

    prompt = f"""
You are the Capital Leverage My Case Agent.

Use this case memory:
{case_data}

User question:
{question}

Answer as a case-focused legal process assistant:
1. What this means for the UMA case
2. Strongest leverage points
3. Weaknesses or risks
4. Evidence/exhibits needed
5. Next steps
6. Draft language if useful

Reminder: Capital Leverage is not a law firm and the user must verify court rules and filings.
"""
    return jsonify(ai_router(prompt, mode))


import base64, pickle
from email.mime.text import MIMEText
from google_auth_oauthlib.flow import Flow
from googleapiclient.discovery import build

GMAIL_SCOPES = [
    "https://www.googleapis.com/auth/gmail.send",
    "https://www.googleapis.com/auth/gmail.readonly"
]

@app.route("/gmail/connect")
def gmail_connect():
    flow = Flow.from_client_secrets_file(
        "credentials.json",
        scopes=GMAIL_SCOPES,
        redirect_uri="https://capitalleverage-1.onrender.com/gmail/callback"
    )
    auth_url, _ = flow.authorization_url(
        access_type="offline",
        include_granted_scopes="true",
        prompt="consent"
    )
    return redirect(auth_url)

@app.route("/gmail/callback")
def gmail_callback():
    flow = Flow.from_client_secrets_file(
        "credentials.json",
        scopes=GMAIL_SCOPES,
        redirect_uri="https://capitalleverage-1.onrender.com/gmail/callback"
    )
    flow.fetch_token(authorization_response=request.url)

    creds = flow.credentials
    with open("gmail_token.pickle", "wb") as token:
        pickle.dump(creds, token)

    return "Gmail connected successfully. You can close this page."

@app.route("/gmail/send", methods=["POST"])
def gmail_send():
    data = request.json
    to = data.get("to")
    subject = data.get("subject")
    body = data.get("body")

    if not to or not subject or not body:
        return jsonify({"success": False, "message": "Missing to, subject, or body"}), 400

    if not os.path.exists("gmail_token.pickle"):
        return jsonify({"success": False, "message": "Gmail not connected yet. Go to /gmail/connect first."}), 400

    with open("gmail_token.pickle", "rb") as token:
        creds = pickle.load(token)

    service = build("gmail", "v1", credentials=creds)

    msg = MIMEText(body)
    msg["to"] = to
    msg["subject"] = subject

    raw = base64.urlsafe_b64encode(msg.as_bytes()).decode()
    sent = service.users().messages().send(userId="me", body={"raw": raw}).execute()

    return jsonify({"success": True, "message": "Email sent.", "id": sent.get("id")})

@app.route("/trading/account", methods=["GET"])
def trading_account():
    import os
    from dotenv import load_dotenv
    from alpaca.trading.client import TradingClient

    load_dotenv(dotenv_path=".env", override=True)

    key = os.getenv("ALPACA_PAPER_API_KEY")
    secret = os.getenv("ALPACA_PAPER_SECRET_KEY")

    if not key or not secret:
        return jsonify({"success": False, "message": "Missing Alpaca paper keys"}), 400

    try:
        client = TradingClient(key, secret, paper=True)
        acct = client.get_account()

        buying_power = float(acct.buying_power) if acct.buying_power is not None else 0.0
        status = str(acct.status.value if hasattr(acct.status, "value") else acct.status)

        return jsonify({
            "success": True,
            "status": status,
            "buying_power": buying_power,
            "trading_mode": str(os.getenv("TRADING_MODE", "paper")),
            "max_stop_loss_percent": float(os.getenv("MAX_STOP_LOSS_PERCENT", "1"))
        })
    except Exception as e:
        return jsonify({"success": False, "message": f"Trading route error: {str(e)}"}), 500

# =========================
# SIMPLE PAPER BOT ROUTES
# =========================

paper_bot_state = {
    "running": False,
    "mode": "paper",
    "strategy": "Capital Leverage 1% Risk Paper Strategy"
}

@app.route("/trading/bot/status", methods=["GET"])
def trading_bot_status():
    return jsonify({
        "success": True,
        "running": paper_bot_state["running"],
        "mode": paper_bot_state["mode"],
        "strategy": paper_bot_state["strategy"]
    })

@app.route("/trading/bot/start", methods=["POST"])
def trading_bot_start():
    paper_bot_state["running"] = True
    return jsonify({
        "success": True,
        "message": "Paper bot started.",
        "running": True
    })

@app.route("/trading/bot/stop", methods=["POST"])
def trading_bot_stop():
    paper_bot_state["running"] = False
    return jsonify({
        "success": True,
        "message": "Paper bot stopped.",
        "running": False
    })



@app.route("/trading/bot/run-strategy", methods=["POST"])
def trading_bot_run_strategy():
    data = request.get_json(silent=True) or {}
    symbol = (data.get("symbol") or data.get("prompt") or "SPY").strip().upper()
    timeframe = data.get("timeframe") or "5Min"

    if not paper_bot_state["running"]:
        return jsonify({
            "success": False,
            "message": "Start the paper bot first, then run strategy."
        }), 400

    try:
        scan = trading_market_scan().get_json()
        watchlist = scan.get("watchlist", [])
    except Exception:
        scan = {}
        watchlist = []

    item = None
    for row in watchlist:
        if row.get("symbol", "").upper() == symbol:
            item = row
            break

    if not item and watchlist:
        item = watchlist[0]
        symbol = item.get("symbol", symbol)

    price = item.get("price") if item else None
    high = item.get("high") if item else None
    low = item.get("low") if item else None
    change_pct = item.get("change_pct") if item else None
    signal = item.get("signal") if item else "Neutral"

    if signal == "Bullish":
        bias = "Bullish above key reclaim/breakout levels"
        entry = f"Watch {symbol} for a break or reclaim near the day high {high}. Confirm with strength before entry."
        stop = f"Keep stop below failed reclaim or under structure near {low}. Risk max 1%."
        target = "Target 2R first, then 3R if momentum holds."
        avoid = "No trade if it rejects the high, loses momentum, or breaks VWAP/support."
    elif signal == "Bearish":
        bias = "Bearish / weak until it reclaims structure"
        entry = f"Watch {symbol} for rejection near resistance or failed reclaim. Paper short only if your broker/setup supports it."
        stop = "Stop above failed breakout/reclaim zone. Risk max 1%."
        target = "Target prior support or 2R. Do not chase after big drop."
        avoid = "No trade if it snaps back above reclaim level with volume."
    else:
        bias = "Neutral / wait for confirmation"
        entry = f"Wait for {symbol} to choose direction. Do not enter in chop."
        stop = "Stop must stay within 1% risk. No setup, no trade."
        target = "Only target after confirmed breakout or breakdown."
        avoid = "Avoid if price stays between high/low without direction."

    return jsonify({
        "success": True,
        "symbol": symbol,
        "timeframe": timeframe,
        "strategy": "Capital Leverage Trade Card — 1% Risk Paper Strategy",
        "price": price,
        "high": high,
        "low": low,
        "change_pct": change_pct,
        "signal": signal,
        "bias": bias,
        "entry_rule": entry,
        "stop_loss_rule": stop,
        "take_profit_rule": target,
        "no_trade_rule": avoid,
        "message": f"Trade card built for {symbol} using live market scan."
    })

@app.route("/trading/market-scan", methods=["GET"])
def trading_market_scan():
    import os, requests
    from dotenv import load_dotenv

    load_dotenv(dotenv_path=".env", override=True)

    key = os.getenv("ALPACA_PAPER_API_KEY")
    secret = os.getenv("ALPACA_PAPER_SECRET_KEY")

    symbols = ["SPY", "QQQ", "AAPL", "TSLA", "NVDA", "AMD", "MSFT", "META"]

    if not key or not secret:
        return jsonify({
            "success": False,
            "message": "Missing Alpaca paper API keys in .env"
        }), 400

    url = "https://data.alpaca.markets/v2/stocks/snapshots"
    headers = {
        "APCA-API-KEY-ID": key,
        "APCA-API-SECRET-KEY": secret
    }

    try:
        r = requests.get(url, headers=headers, params={"symbols": ",".join(symbols)}, timeout=30)

        if r.status_code != 200:
            return jsonify({
                "success": False,
                "message": f"Alpaca market scan failed: {r.status_code} {r.text[:300]}"
            }), 500

        data = r.json()
        snapshots = data.get("snapshots", data)

        rows = []

        for sym in symbols:
            snap = snapshots.get(sym, {}) or {}
            trade = snap.get("latestTrade") or {}
            day = snap.get("dailyBar") or {}

            price = trade.get("p") or day.get("c")
            high = day.get("h")
            low = day.get("l")
            open_price = day.get("o")

            change_pct = None
            signal = "Neutral"

            if price and open_price:
                change_pct = round(((price - open_price) / open_price) * 100, 2)
                if change_pct >= 1:
                    signal = "Bullish"
                elif change_pct <= -1:
                    signal = "Bearish"

            rows.append({
                "symbol": sym,
                "price": price,
                "high": high,
                "low": low,
                "open": open_price,
                "change_pct": change_pct,
                "signal": signal
            })

        rows.append({
            "symbol": "BTC/USD",
            "price": None,
            "high": None,
            "low": None,
            "change_pct": None,
            "signal": "Crypto research mode"
        })

        rows.append({
            "symbol": "GBPJPY",
            "price": None,
            "high": None,
            "low": None,
            "change_pct": None,
            "signal": "OpenAlice forex research mode"
        })

        top = sorted(
            [x for x in rows if isinstance(x.get("change_pct"), (int, float))],
            key=lambda x: x["change_pct"],
            reverse=True
        )[:3]

        return jsonify({
            "success": True,
            "mode": "paper",
            "risk_rule": "1% max stop loss",
            "watchlist": rows,
            "top_3": top,
            "summary": "Here is the move. Capital Leverage scanned the board. Focus on the strongest names, stay in paper mode, and keep risk capped at 1%."
        })

    except Exception as e:
        return jsonify({
            "success": False,
            "message": str(e)
        }), 500


@app.route("/trading/smart-command", methods=["POST"])
def trading_smart_command():
    data = request.get_json(silent=True) or {}
    command = data.get("command", "daily")
    symbol = data.get("symbol", "SPY")

    try:
        scan = trading_market_scan().get_json()
    except Exception as e:
        scan = {"watchlist": [], "top_3": [], "message": str(e)}

    prompt = f"""
You are Capital Leverage Trading Desk.

Selected symbol: {symbol}
Command: {command}

Live market scan:
{scan}

Rules:
- Paper trading only.
- No profit guarantees.
- Max stop loss is 1% per trade.
- Use the live scan to choose the best opportunity.
- Give market-aware guidance.

Return:
HERE IS THE MOVE
BEST OPPORTUNITY RIGHT NOW
TOP 3 TO WATCH
NO-TRADE / AVOID LIST
ENTRY TRIGGER
STOP LOSS PLAN
TAKE PROFIT PLAN
WHAT MARKET/NEWS INFO MATTERS
NEXT ACTION
"""

    try:
        return jsonify(ai_router(prompt, "research"))
    except Exception as e:
        return jsonify({"answer": f"Smart trading command failed: {e}", "provider": "System"}), 500


# =========================
# MY CASE FILE VAULT
# =========================
MYCASE_UPLOAD_DIR = UPLOAD_DIR / "mycase"
MYCASE_TEXT_DIR = TEXT_DIR / "mycase"

MYCASE_UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
MYCASE_TEXT_DIR.mkdir(parents=True, exist_ok=True)

MYCASE_DOCS_FILE = DATA_DIR / "mycase_documents.json"

@app.route("/mycase/files", methods=["GET"])
def get_mycase_files():
    return jsonify(load_json(MYCASE_DOCS_FILE))

@app.route("/mycase/upload", methods=["POST"])
def upload_mycase_file():
    if "file" not in request.files:
        return jsonify({"error": "No file uploaded"}), 400

    f = request.files["file"]
    filename = secure_filename(f.filename)

    if not filename:
        return jsonify({"error": "Missing filename"}), 400

    save_path = MYCASE_UPLOAD_DIR / filename
    f.save(save_path)

    try:
        extracted_text = extract_text_from_file(save_path)
    except Exception as e:
        extracted_text = f"[Text extraction failed: {e}]"

    text_file = MYCASE_TEXT_DIR / f"{filename}.txt"
    text_file.write_text(extracted_text)

    doc = {
        "id": str(int(datetime.now().timestamp())),
        "filename": filename,
        "path": str(save_path),
        "text_path": str(text_file),
        "text_length": len(extracted_text),
        "created_at": datetime.now().isoformat()
    }

    docs = load_json(MYCASE_DOCS_FILE)
    docs.append(doc)
    save_json(MYCASE_DOCS_FILE, docs)

    return jsonify(doc)

@app.route("/mycase/analyze", methods=["POST"])
def analyze_mycase():
    data = request.get_json(silent=True) or {}
    question = (data.get("question") or "").strip()

    docs = load_json(MYCASE_DOCS_FILE)

    if not docs:
        return jsonify({
            "answer": "No My Case files uploaded yet. Upload your UMA exhibits, motion papers, screenshots, or evidence first.",
            "provider": "System"
        })

    combined = []
    for d in docs[-12:]:
        text_path = d.get("text_path")
        if text_path and Path(text_path).exists():
            try:
                txt = Path(text_path).read_text(errors="ignore")
                combined.append(f"\n\n===== FILE: {d.get('filename')} =====\n{txt[:25000]}")
            except Exception:
                pass

    evidence_text = "\n".join(combined)[:120000]

    if not question:
        question = "Review my uploaded UMA case files and tell me the strongest evidence, contradictions, damages proof, discovery targets, and next legal moves."

    prompt = f"""
You are Capital Leverage My Case Agent.

This is the user's Harrison v. UMA / Hire Image case command center.

Use the uploaded case evidence below to answer the user's question.
Focus on:
- strongest facts
- timeline issues
- contradictions
- FCRA / fraud / unjust enrichment / damages angles
- missing evidence
- discovery targets
- settlement leverage
- next step recommendations

User question:
{question}

Uploaded case evidence:
{evidence_text}

Return:
1. Direct answer
2. Strongest evidence from uploaded files
3. Weak points / gaps
4. What to do next
"""

    return jsonify(ai_router(prompt, "document"))
# =========================
# OPENALICE STATUS ROUTES
# =========================

@app.route("/trading/openalice/ping", methods=["GET"])
def trading_openalice_ping():
    import os
    from dotenv import load_dotenv
    load_dotenv(dotenv_path=".env", override=True)

    return jsonify({
        "success": True,
        "openalice_base_url": os.getenv("OPENALICE_BASE_URL", "http://127.0.0.1:47331"),
        "token_present": bool(os.getenv("OPENALICE_ADMIN_TOKEN"))
    })

@app.route("/trading/openalice/account-check", methods=["GET"])
def trading_openalice_account_check():
    import os, requests
    from dotenv import load_dotenv
    load_dotenv(dotenv_path=".env", override=True)

    base = os.getenv("OPENALICE_BASE_URL", "http://127.0.0.1:47331").rstrip("/")
    token = os.getenv("OPENALICE_ADMIN_TOKEN", "").strip()

    if not token:
        return jsonify({"success": False, "message": "Missing OPENALICE_ADMIN_TOKEN"}), 400

    sess = requests.Session()
    login = sess.post(f"{base}/api/auth/login", json={"token": token}, timeout=30)

    if login.status_code != 200:
        return jsonify({
            "success": False,
            "message": f"OpenAlice login failed: {login.status_code} {login.text}"
        }), 500

    status = sess.get(f"{base}/api/auth/status", timeout=30)

    return jsonify({
        "success": True,
        "message": "OpenAlice is connected.",
        "openalice_status": status.json() if status.headers.get("content-type", "").startswith("application/json") else status.text
    })
# =========================
# OPENALICE STRATEGY ROUTES
# =========================

def openalice_login_session():
    import os, requests
    from dotenv import load_dotenv
    load_dotenv(dotenv_path=".env", override=True)

    base = os.getenv("OPENALICE_BASE_URL", "http://127.0.0.1:47331").rstrip("/")
    token = os.getenv("OPENALICE_ADMIN_TOKEN", "").strip()

    sess = requests.Session()
    login = sess.post(f"{base}/api/auth/login", json={"token": token}, timeout=30)

    if login.status_code != 200:
        return None, base, f"OpenAlice login failed: {login.status_code} {login.text}"

    return sess, base, None


@app.route("/trading/openalice/quick-chat", methods=["POST"])
def trading_openalice_quick_chat():
    data = request.get_json(silent=True) or {}
    prompt = data.get("prompt", "Build a Capital Leverage trading plan.")

    sess, base, err = openalice_login_session()
    if err:
        return jsonify({"success": False, "message": err}), 500

    full_prompt = f"""
You are Capital Leverage Trading Desk.

Use smart trading desk style.
Paper trading only.
Max stop loss: 1%.
No profit guarantees.

User request:
{prompt}

Return:
HERE IS THE MOVE
BEST SETUP
TOP 3 WATCHLIST
ENTRY TRIGGER
STOP LOSS PLAN
TAKE PROFIT PLAN
NO-TRADE LIST
NEXT ACTION
"""

    r = sess.post(
        f"{base}/api/workspaces/quick-chat",
        json={"prompt": full_prompt},
        timeout=120
    )

    try:
        body = r.json()
    except Exception:
        body = r.text

    return jsonify({
        "success": r.status_code in [200, 201, 202],
        "status_code": r.status_code,
        "response": body
    })


@app.route("/trading/openalice/strategy", methods=["POST"])
def trading_openalice_strategy():
    data = request.get_json(silent=True) or {}
    symbol = data.get("symbol") or data.get("prompt") or "SPY"

    try:
        scan = trading_market_scan().get_json()
    except Exception:
        scan = {}

    prompt = f"""
Build a market-aware strategy for {symbol}.

Live Capital Leverage market scan:
{scan}

Rules:
Paper mode only.
Risk max 1%.
Use the scan.
Give the best opportunity, levels, entry, stop, target, invalidation, and next action.
"""

    return trading_openalice_quick_chat.__wrapped__() if False else _openalice_strategy_post(prompt)


def _openalice_strategy_post(prompt):
    sess, base, err = openalice_login_session()
    if err:
        return jsonify({"success": False, "message": err}), 500

    r = sess.post(
        f"{base}/api/workspaces/quick-chat",
        json={"prompt": prompt},
        timeout=120
    )

    try:
        body = r.json()
    except Exception:
        body = r.text

    return jsonify({
        "success": r.status_code in [200, 201, 202],
        "status_code": r.status_code,
        "answer": body
    })
# =========================
# AUTO PAPER TRADE
# =========================

@app.route("/trading/bot/auto-paper-trade", methods=["POST"])
def auto_paper_trade():
    import os
    from dotenv import load_dotenv
    from alpaca.trading.client import TradingClient
    from alpaca.trading.requests import MarketOrderRequest, TakeProfitRequest, StopLossRequest
    from alpaca.trading.enums import OrderSide, TimeInForce, OrderClass

    load_dotenv(dotenv_path=".env", override=True)

    key = os.getenv("ALPACA_PAPER_API_KEY")
    secret = os.getenv("ALPACA_PAPER_SECRET_KEY")

    if not key or not secret:
        return jsonify({"success": False, "message": "Missing Alpaca paper keys."}), 400

    scan = trading_market_scan().get_json()
    top = scan.get("top_3", [])

    pick = None
    for item in top:
        if item.get("signal") == "Bullish" and item.get("price"):
            pick = item
            break

    if not pick:
        return jsonify({
            "success": False,
            "message": "No bullish paper-trade setup found right now. Bot will not force a trade."
        }), 400

    symbol = pick["symbol"]
    price = float(pick["price"])
    stop_price = round(price * 0.99, 2)
    take_profit_price = round(price * 1.02, 2)

    client = TradingClient(key, secret, paper=True)

    order = MarketOrderRequest(
        symbol=symbol,
        qty=1,
        side=OrderSide.BUY,
        time_in_force=TimeInForce.DAY,
        order_class=OrderClass.BRACKET,
        take_profit=TakeProfitRequest(limit_price=take_profit_price),
        stop_loss=StopLossRequest(stop_price=stop_price)
    )

    submitted = client.submit_order(order)

    return jsonify({
        "success": True,
        "message": "Paper trade placed.",
        "symbol": symbol,
        "qty": 1,
        "estimated_entry": price,
        "stop_loss": stop_price,
        "take_profit": take_profit_price,
        "risk_rule": "1% stop loss",
        "order_id": str(submitted.id)
    })
# =========================
# MUSIC MANAGER OS
# =========================

@app.route("/business/music-manager", methods=["POST"])
def business_music_manager():
    data = request.get_json(silent=True) or {}
    task = data.get("task", "build")
    details = data.get("details", "")

    prompt = f"""
You are Capital Leverage Business Agent.

Build and advise on Music Manager OS.

Vision:
Music Manager OS is a multi-user artist growth and management platform.
It acts like a real manager for independent artists.

Core parts:
1. Manager View - daily tasks and growth advice
2. Tracker - clicks, views, followers, streams, campaign data
3. Campaign Tool - promo links, goals, deadlines, traffic sources
4. Client Portal - each artist has their own login/dashboard
5. Admin Dashboard - manager can oversee all artists
6. Automation - content ideas, hashtags, repost planning, reports

Task requested:
{task}

User details:
{details}

Return:
1. What this feature should do
2. Screen layout
3. Database/data needed
4. AI behavior
5. Buttons/functions needed
6. First version build plan
7. Next action
"""

    return jsonify(ai_router(prompt, "business"))
# =========================
# MUSIC MANAGER OS DATA
# =========================

MUSIC_ARTISTS_FILE = DATA_DIR / "music_artists.json"
MUSIC_CAMPAIGNS_FILE = DATA_DIR / "music_campaigns.json"
MUSIC_LINKS_FILE = DATA_DIR / "music_links.json"
MUSIC_TASKS_FILE = DATA_DIR / "music_tasks.json"

@app.route("/music/artists", methods=["GET", "POST"])
def music_artists():
    if request.method == "GET":
        return jsonify(load_json(MUSIC_ARTISTS_FILE))

    data = request.get_json(silent=True) or {}
    artist = {
        "id": str(int(datetime.now().timestamp())),
        "name": data.get("name", "Untitled Artist"),
        "genre": data.get("genre", ""),
        "goal": data.get("goal", ""),
        "created_at": datetime.now().isoformat()
    }

    artists = load_json(MUSIC_ARTISTS_FILE)
    artists.append(artist)
    save_json(MUSIC_ARTISTS_FILE, artists)
    return jsonify(artist)

@app.route("/music/campaigns", methods=["GET", "POST"])
def music_campaigns():
    if request.method == "GET":
        return jsonify(load_json(MUSIC_CAMPAIGNS_FILE))

    data = request.get_json(silent=True) or {}
    campaign = {
        "id": str(int(datetime.now().timestamp())),
        "artist": data.get("artist", ""),
        "title": data.get("title", "Untitled Campaign"),
        "platform": data.get("platform", "TikTok / YouTube"),
        "goal": data.get("goal", ""),
        "status": data.get("status", "Active"),
        "created_at": datetime.now().isoformat()
    }

    campaigns = load_json(MUSIC_CAMPAIGNS_FILE)
    campaigns.append(campaign)
    save_json(MUSIC_CAMPAIGNS_FILE, campaigns)
    return jsonify(campaign)

@app.route("/music/links", methods=["GET", "POST"])
def music_links():
    if request.method == "GET":
        return jsonify(load_json(MUSIC_LINKS_FILE))

    data = request.get_json(silent=True) or {}
    link = {
        "id": str(int(datetime.now().timestamp())),
        "artist": data.get("artist", ""),
        "campaign": data.get("campaign", ""),
        "url": data.get("url", ""),
        "clicks": 0,
        "created_at": datetime.now().isoformat()
    }

    links = load_json(MUSIC_LINKS_FILE)
    links.append(link)
    save_json(MUSIC_LINKS_FILE, links)
    return jsonify(link)


@app.route("/music/manager-agent", methods=["POST"])
def music_manager_agent():
    data = request.get_json(silent=True) or {}
    task = data.get("task", "daily")
    artist = data.get("artist", "")
    genre = data.get("genre", "")
    details = data.get("details", "")

    artists = load_json(MUSIC_ARTISTS_FILE)
    campaigns = load_json(MUSIC_CAMPAIGNS_FILE)
    links = load_json(MUSIC_LINKS_FILE)

    prompt = f"""
You are Music Manager OS inside Capital Leverage.

You are not a normal chatbot.
You are the artist's manager.

Your job:
- Tell the artist exactly what to do today.
- Build campaigns.
- Write titles, captions, hashtags, and posting plans.
- Create video ideas.
- Help the artist become known.
- Think like a real manager trying to grow views, fans, comments, and momentum.

Task:
{task}

Artist:
{artist}

Genre:
{genre}

Details:
{details}

Saved artists:
{artists}

Saved campaigns:
{campaigns}

Saved links:
{links}

Response style:
Start with: HERE IS THE MOVE.
Be confident.
Give direct steps.
Do not sound generic.

Return:
1. HERE IS THE MOVE
2. What to post today
3. 10 content/video ideas if relevant
4. Titles/captions/hashtags if relevant
5. YouTube/TikTok plan
6. Campaign move
7. What needs attention
8. Next 3 actions
"""

    return jsonify(ai_router(prompt, "business"))
# =========================
# MUSIC MANAGER AUTO CAMPAIGN AGENT
# =========================

MUSIC_AUTOMATION_FILE = DATA_DIR / "music_automation_tasks.json"

# =========================
# MUSIC MANAGER AUTO CAMPAIGN AGENT
# =========================

MUSIC_AUTOMATION_FILE = DATA_DIR / "music_automation_tasks.json"

@app.route("/music/auto-campaign", methods=["POST"])
def music_auto_campaign():
    data = request.get_json(silent=True) or {}

    artist = data.get("artist", "")
    genre = data.get("genre", "")
    song = data.get("song", "")
    goal = data.get("goal", "Grow views, engagement, followers, and campaign momentum.")
    platforms = data.get("platforms", "YouTube Shorts, TikTok, Instagram Reels")

    prompt = f"""
You are Music Manager OS inside Capital Leverage.

You are the artist's campaign manager.
Your job is to automatically build the full campaign like a real manager.

Artist: {artist}
Genre: {genre}
Song/project: {song}
Goal: {goal}
Platforms: {platforms}

Return:
1. HERE IS THE MOVE
2. Campaign Name
3. 7-Day Campaign Schedule
4. 10 Video Ideas
5. 10 Titles
6. 10 Captions
7. Hashtag Sets
8. Posting Times
9. YouTube Shorts Plan
10. TikTok Plan
11. Repost Strategy
12. Fan Engagement Tasks
13. Daily Checklist
14. What the agent should do next automatically
"""

    result = ai_router(prompt, "business")
    answer = result.get("answer", str(result)) if isinstance(result, dict) else str(result)

    task = {
        "id": str(int(datetime.now().timestamp())),
        "artist": artist,
        "genre": genre,
        "song": song,
        "goal": goal,
        "platforms": platforms,
        "status": "Campaign Plan Built",
        "campaign_output": answer,
        "created_at": datetime.now().isoformat()
    }

    tasks = load_json(MUSIC_AUTOMATION_FILE)
    tasks.append(task)
    save_json(MUSIC_AUTOMATION_FILE, tasks)

    return jsonify({
        "success": True,
        "message": "Music Manager Agent built the campaign.",
        "task": task,
        "answer": answer
    })

@app.route("/music/automation-tasks", methods=["GET"])
def music_automation_tasks():
    return jsonify(load_json(MUSIC_AUTOMATION_FILE))

if __name__ == "__main__":
    app.run(host="0.0.0.0", port=9000, debug=True)
