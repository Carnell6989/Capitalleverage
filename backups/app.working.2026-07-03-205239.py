from flask import Flask, render_template, request, jsonify
from ai.router import ai_router
import json
from pathlib import Path
from datetime import datetime
from werkzeug.utils import secure_filename
import fitz
import docx

app = Flask(__name__)

DATA_DIR = Path("data")
UPLOAD_DIR = Path("uploads")
TEXT_DIR = Path("data/extracted_text")

DATA_DIR.mkdir(exist_ok=True)
UPLOAD_DIR.mkdir(exist_ok=True)
TEXT_DIR.mkdir(exist_ok=True)

CASES_FILE = DATA_DIR / "cases.json"
DOCS_FILE = DATA_DIR / "documents.json"

def load_json(path):
    if not path.exists():
        return []
    return json.loads(path.read_text())

def save_json(path, data):
    path.write_text(json.dumps(data, indent=2))

def extract_text_from_file(path):
    suffix = path.suffix.lower()

    if suffix == ".pdf":
        text = ""
        with fitz.open(path) as pdf:
            for i, page in enumerate(pdf):
                text += f"\n\n--- Page {i+1} ---\n"
                text += page.get_text()
        return text.strip()

    if suffix == ".docx":
        document = docx.Document(path)
        return "\n".join([p.text for p in document.paragraphs]).strip()

    return path.read_text(errors="ignore").strip()

@app.route("/")
def home():
    return render_template("index.html")

@app.route("/ask", methods=["POST"])
def ask():
    message = request.json.get("message", "").strip()
    mode = request.json.get("mode", "fast")

    if not message:
        return jsonify({"answer": "Type a message first.", "provider": "System"})

    return jsonify(ai_router(message, mode))

@app.route("/cases", methods=["GET"])
def get_cases():
    return jsonify(load_json(CASES_FILE))

@app.route("/cases", methods=["POST"])
def create_case():
    data = request.json
    case = {
        "id": str(int(datetime.now().timestamp())),
        "name": data.get("name", "Untitled Case"),
        "type": data.get("type", "Legal"),
        "status": data.get("status", "Active"),
        "notes": data.get("notes", ""),
        "created_at": datetime.now().isoformat()
    }

    cases = load_json(CASES_FILE)
    cases.append(case)
    save_json(CASES_FILE, cases)
    return jsonify(case)

@app.route("/documents", methods=["GET"])
def get_documents():
    return jsonify(load_json(DOCS_FILE))

@app.route("/documents/upload", methods=["POST"])
def upload_document():
    if "file" not in request.files:
        return jsonify({"error": "No file uploaded"}), 400

    f = request.files["file"]
    filename = secure_filename(f.filename)

    if not filename:
        return jsonify({"error": "Missing filename"}), 400

    save_path = UPLOAD_DIR / filename
    f.save(save_path)

    try:
        extracted_text = extract_text_from_file(save_path)
    except Exception as e:
        extracted_text = f"[Text extraction failed: {e}]"

    text_file = TEXT_DIR / f"{filename}.txt"
    text_file.write_text(extracted_text)

    doc = {
        "id": str(int(datetime.now().timestamp())),
        "filename": filename,
        "path": str(save_path),
        "text_path": str(text_file),
        "text_length": len(extracted_text),
        "created_at": datetime.now().isoformat()
    }

    docs = load_json(DOCS_FILE)
    docs.append(doc)
    save_json(DOCS_FILE, docs)

    return jsonify(doc)


@app.route("/documents/delete", methods=["POST"])
def delete_document():
    data = request.json
    filename = data.get("filename")

    docs = load_json(DOCS_FILE)
    docs = [d for d in docs if d.get("filename") != filename]
    save_json(DOCS_FILE, docs)

    for folder in [UPLOAD_DIR, TEXT_DIR]:
        for file in folder.glob(f"{filename}*"):
            try:
                file.unlink()
            except Exception:
                pass

    return jsonify({"status": "deleted", "filename": filename})

@app.route("/documents/analyze", methods=["POST"])
def analyze_document():
    data = request.json
    filename = data.get("filename")
    question = data.get("question", "Summarize this document.")

    text_file = TEXT_DIR / f"{filename}.txt"

    if not text_file.exists():
        return jsonify({"answer": "Extracted document text not found. Re-upload the file.", "provider": "System"})

    text = text_file.read_text(errors="ignore")

    if not text.strip():
        return jsonify({"answer": "No readable text found in this document. OCR will be needed next.", "provider": "System"})

    prompt = f"""
You are analyzing an uploaded legal/business document.

Filename: {filename}

User question:
{question}

Document text:
{text[:30000]}

Return:
1. Short summary
2. Important dates
3. Important people/entities
4. Main claims/arguments
5. Weaknesses or contradictions
6. Recommended next steps
"""

    return jsonify(ai_router(prompt, "document"))


@app.route("/templates/generate", methods=["POST"])
def generate_template():
    data = request.json
    template_type = data.get("template_type", "")
    facts = data.get("facts", "")

    prompt = f"""
Create a professional Capital Leverage document.

Document type:
{template_type}

User facts:
{facts}

Rules:
- Use clear headings.
- Include placeholders for missing names, dates, addresses, account numbers, and signatures.
- Make it copy-paste ready.
- Include a short disclaimer if legal/credit/housing related.
"""

    return jsonify(ai_router(prompt, "document"))


OPPS_FILE = DATA_DIR / "opportunities.json"

@app.route("/opportunities", methods=["GET"])
def get_opportunities():
    return jsonify(load_json(OPPS_FILE))

@app.route("/opportunities", methods=["POST"])
def create_opportunity():
    data = request.json
    opp = {
        "id": str(int(datetime.now().timestamp())),
        "title": data.get("title", "Untitled Opportunity"),
        "category": data.get("category", "Housing"),
        "status": data.get("status", "Researching"),
        "deadline": data.get("deadline", ""),
        "notes": data.get("notes", ""),
        "created_at": datetime.now().isoformat()
    }

    opps = load_json(OPPS_FILE)
    opps.append(opp)
    save_json(OPPS_FILE, opps)
    return jsonify(opp)

@app.route("/opportunities/plan", methods=["POST"])
def opportunity_plan():
    data = request.json
    goal = data.get("goal", "")

    prompt = f"""
Create a Capital Leverage opportunity plan.

User goal:
{goal}

Focus on:
1. Housing opportunities
2. Credit repair opportunities
3. Business funding opportunities
4. Legal/document leverage opportunities
5. Applications to look for
6. Documents needed
7. 7-day action plan
8. 30-day action plan

Important:
If current program availability is needed, say live web search must verify it.
"""

    return jsonify(ai_router(prompt, "document"))


EVIDENCE_FILE = DATA_DIR / "evidence.json"

@app.route("/evidence", methods=["GET"])
def get_evidence():
    return jsonify(load_json(EVIDENCE_FILE))

@app.route("/evidence", methods=["POST"])
def create_evidence():
    data = request.json

    evidence = {
        "id": str(int(datetime.now().timestamp())),
        "title": data.get("title", "Untitled Evidence"),
        "category": data.get("category", "General"),
        "date": data.get("date", ""),
        "case_name": data.get("case_name", ""),
        "notes": data.get("notes", ""),
        "created_at": datetime.now().isoformat()
    }

    items = load_json(EVIDENCE_FILE)
    items.append(evidence)
    save_json(EVIDENCE_FILE, items)

    return jsonify(evidence)


@app.route("/email/intake", methods=["POST"])
def email_intake():
    data = request.json
    sender = data.get("sender", "")
    subject = data.get("subject", "")
    body = data.get("body", "")

    prompt = f"""
Review this email for Capital Leverage.

From:
{sender}

Subject:
{subject}

Email:
{body}

Return:
1. Summary
2. Category: legal, credit, housing, business, personal, or urgent
3. What it means
4. Leverage opportunities
5. Recommended next steps
6. Draft a professional reply
7. If useful, explain how this could become evidence
"""
    return jsonify(ai_router(prompt, "document"))


MEMORY_FILE = DATA_DIR / "memory.json"

@app.route("/memory", methods=["GET"])
def get_memory():
    memory = load_json(MEMORY_FILE)
    if isinstance(memory, list):
        return jsonify({"memory": ""})
    return jsonify(memory)

@app.route("/memory", methods=["POST"])
def save_memory():
    data = request.json
    memory = {
        "memory": data.get("memory", ""),
        "updated_at": datetime.now().isoformat()
    }
    save_json(MEMORY_FILE, memory)
    return jsonify(memory)

@app.route("/health")
def health():
    return jsonify({"status": "ok", "app": "Capital Leverage 2.5"})


@app.route("/my-case/all-agents", methods=["POST"])
def my_case_all_agents():
    data = request.json
    question = data.get("question", "")
    case_data = load_json(DATA_DIR / "my_case.json")

    prompt = f"""
You are the Capital Leverage Multi-Agent Case Team.

Case Memory:
{case_data}

User Question:
{question}

Work as a team:
1. Legal Agent
2. Credit/FCRA Agent
3. Document Evidence Agent
4. Damages Agent
5. Settlement/Discovery Agent

Return:
1. Legal Agent analysis
2. Credit/FCRA Agent analysis
3. Document/Evidence Agent analysis
4. Damages Agent analysis
5. Settlement/Discovery Agent analysis
6. Combined best strategy
7. Next 5 actions
8. Draft language if useful

Reminder: not legal advice; verify rules, deadlines, and filings.
"""
    return jsonify(ai_router(prompt, "backup"))



from ai.web_research import web_research

@app.route("/my-case/research", methods=["POST"])
def my_case_research():
    data = request.json
    question = data.get("question", "")
    case_data = load_json(DATA_DIR / "my_case.json")

    research = web_research(question, max_results=5)

    prompt = f"""
You are the Capital Leverage My Case Web Research Agent.

Case Memory:
{case_data}

User Question:
{question}

Live Web Research Results:
{research}

Answer using the research above. Do not invent citations.
Return:
1. What the research says
2. How it applies to the UMA case
3. Strongest leverage points
4. Risks or weaknesses
5. Next actions
6. Sources used with URLs
"""
    return jsonify(ai_router(prompt, "backup"))



@app.route("/my-case/task", methods=["POST"])
def my_case_task():
    data = request.json
    task = data.get("task", "case_strategy")
    question = data.get("question", "")
    case_data = load_json(DATA_DIR / "my_case.json")

    task_prompts = {
        "legal_review": "Run a legal review of the UMA opposition and surviving claims. Identify strongest arguments, weak points, missing support, and next steps.",
        "case_law": "Use web research to find relevant law, cases, statutes, and rules for the user question. Explain how each source may help the UMA case. Do not invent citations.",
        "discovery": "Build discovery requests for UMA: interrogatories, requests for production, admissions, deposition topics, and documents to demand.",
        "cfpb": "Draft a CFPB complaint connected to the facts. Include company, product/service, issue, timeline, harm, requested resolution, and documents to attach.",
        "damages": "Build a damages summary: Pell LEU burn, debt/interest, reputational harm, emotional distress, opportunity loss, evidence needed, and expert proof.",
        "settlement": "Create a settlement leverage package: liability theory, damages theory, evidence list, negotiation posture, demand structure, and next actions.",
        "timeline": "Build a clean case timeline from the facts and identify deadlines, proof gaps, and exhibit connections.",
        "exhibits": "Create an exhibit/evidence plan: what each exhibit proves, how to label it, where it supports each count, and what is missing."
    }

    task_instruction = task_prompts.get(task, task_prompts["legal_review"])

    research = ""
    if task in ["case_law", "legal_review", "settlement", "cfpb"]:
        try:
            research = web_research(question or task_instruction, max_results=5)
        except Exception as e:
            research = f"Web research unavailable: {e}"

    prompt = f"""
You are the Capital Leverage My Case Command Center.

Case Memory:
{case_data}

Task:
{task_instruction}

User Extra Details / Question:
{question}

Live Web Research if available:
{research}

Return a powerful, case-focused result with:
1. Short summary
2. Key leverage points
3. Evidence needed
4. Draft language if useful
5. Next 5 actions
6. Sources used if web research was provided

Do not invent facts. If current law/rules are needed, say they must be verified.
Capital Leverage is not a law firm.
"""
    return jsonify(ai_router(prompt, "research" if research else "legal"))
@app.route("/my-case", methods=["GET"])
def get_my_case():
    return jsonify(load_json(DATA_DIR / "my_case.json"))

@app.route("/my-case/ask", methods=["POST"])
def ask_my_case():
    data = request.json
    question = data.get("question", "")
    mode = data.get("mode", "legal")
    case_data = load_json(DATA_DIR / "my_case.json")

    prompt = f"""
You are the Capital Leverage My Case Agent.

Use this case memory:
{case_data}

User question:
{question}

Answer as a case-focused legal process assistant:
1. What this means for the UMA case
2. Strongest leverage points
3. Weaknesses or risks
4. Evidence/exhibits needed
5. Next steps
6. Draft language if useful

Reminder: Capital Leverage is not a law firm and the user must verify court rules and filings.
"""
    return jsonify(ai_router(prompt, mode))


import base64, pickle
from email.mime.text import MIMEText
from google_auth_oauthlib.flow import Flow
from googleapiclient.discovery import build

GMAIL_SCOPES = [
    "https://www.googleapis.com/auth/gmail.send",
    "https://www.googleapis.com/auth/gmail.readonly"
]

@app.route("/gmail/connect")
def gmail_connect():
    flow = Flow.from_client_secrets_file(
        "credentials.json",
        scopes=GMAIL_SCOPES,
        redirect_uri="https://capitalleverage-1.onrender.com/gmail/callback"
    )
    auth_url, _ = flow.authorization_url(
        access_type="offline",
        include_granted_scopes="true",
        prompt="consent"
    )
    return redirect(auth_url)

@app.route("/gmail/callback")
def gmail_callback():
    flow = Flow.from_client_secrets_file(
        "credentials.json",
        scopes=GMAIL_SCOPES,
        redirect_uri="https://capitalleverage-1.onrender.com/gmail/callback"
    )
    flow.fetch_token(authorization_response=request.url)

    creds = flow.credentials
    with open("gmail_token.pickle", "wb") as token:
        pickle.dump(creds, token)

    return "Gmail connected successfully. You can close this page."

@app.route("/gmail/send", methods=["POST"])
def gmail_send():
    data = request.json
    to = data.get("to")
    subject = data.get("subject")
    body = data.get("body")

    if not to or not subject or not body:
        return jsonify({"success": False, "message": "Missing to, subject, or body"}), 400

    if not os.path.exists("gmail_token.pickle"):
        return jsonify({"success": False, "message": "Gmail not connected yet. Go to /gmail/connect first."}), 400

    with open("gmail_token.pickle", "rb") as token:
        creds = pickle.load(token)

    service = build("gmail", "v1", credentials=creds)

    msg = MIMEText(body)
    msg["to"] = to
    msg["subject"] = subject

    raw = base64.urlsafe_b64encode(msg.as_bytes()).decode()
    sent = service.users().messages().send(userId="me", body={"raw": raw}).execute()

    return jsonify({"success": True, "message": "Email sent.", "id": sent.get("id")})

@app.route("/trading/account", methods=["GET"])
def trading_account():
    import os
    from dotenv import load_dotenv
    from alpaca.trading.client import TradingClient

    load_dotenv(dotenv_path=".env", override=True)

    key = os.getenv("ALPACA_PAPER_API_KEY")
    secret = os.getenv("ALPACA_PAPER_SECRET_KEY")

    if not key or not secret:
        return jsonify({"success": False, "message": "Missing Alpaca paper keys"}), 400

    try:
        client = TradingClient(key, secret, paper=True)
        acct = client.get_account()

        buying_power = float(acct.buying_power) if acct.buying_power is not None else 0.0
        status = str(acct.status.value if hasattr(acct.status, "value") else acct.status)

        return jsonify({
            "success": True,
            "status": status,
            "buying_power": buying_power,
            "trading_mode": str(os.getenv("TRADING_MODE", "paper")),
            "max_stop_loss_percent": float(os.getenv("MAX_STOP_LOSS_PERCENT", "1"))
        })
    except Exception as e:
        return jsonify({"success": False, "message": f"Trading route error: {str(e)}"}), 500

# =========================
# SIMPLE PAPER BOT ROUTES
# =========================

paper_bot_state = {
    "running": False,
    "mode": "paper",
    "strategy": "Capital Leverage 1% Risk Paper Strategy"
}

@app.route("/trading/bot/status", methods=["GET"])
def trading_bot_status():
    return jsonify({
        "success": True,
        "running": paper_bot_state["running"],
        "mode": paper_bot_state["mode"],
        "strategy": paper_bot_state["strategy"]
    })

@app.route("/trading/bot/start", methods=["POST"])
def trading_bot_start():
    paper_bot_state["running"] = True
    return jsonify({
        "success": True,
        "message": "Paper bot started.",
        "running": True
    })

@app.route("/trading/bot/stop", methods=["POST"])
def trading_bot_stop():
    paper_bot_state["running"] = False
    return jsonify({
        "success": True,
        "message": "Paper bot stopped.",
        "running": False
    })


@app.route("/trading/bot/run-strategy", methods=["POST"])
def trading_bot_run_strategy():
    data = request.get_json(silent=True) or {}
    symbol = data.get("symbol") or data.get("prompt") or "SPY"
    timeframe = data.get("timeframe") or "5Min"

    if not paper_bot_state["running"]:
        return jsonify({
            "success": False,
            "message": "Start the paper bot first, then run strategy."
        }), 400

    return jsonify({
        "success": True,
        "symbol": symbol,
        "timeframe": timeframe,
        "strategy": "Capital Leverage 1% Risk Paper Strategy",
        "entry_rule": f"Watch {symbol} for confirmation. Enter only after a clean breakout, reclaim, or retest. Do not chase.",
        "stop_loss_rule": "Maximum risk is 1% per trade. If the setup breaks structure, exit.",
        "take_profit_rule": "Target 2R to 3R when possible. Take partial profit if momentum slows.",
        "message": f"Paper strategy ran for {symbol} on {timeframe}."
    })


# =========================
# TRADING DESK 2.0 MARKET SCAN
# =========================

@app.route("/trading/market-scan", methods=["GET"])
def trading_market_scan():
    import os, requests
    from dotenv import load_dotenv

    load_dotenv(dotenv_path=".env", override=True)

    key = os.getenv("ALPACA_PAPER_API_KEY")
    secret = os.getenv("ALPACA_PAPER_SECRET_KEY")

    symbols = ["SPY", "QQQ", "AAPL", "TSLA", "NVDA", "AMD", "MSFT", "META"]

    if not key or not secret:
        return jsonify({
            "success": False,
            "message": "Missing Alpaca paper API keys in .env"
        }), 400

    url = "https://data.alpaca.markets/v2/stocks/snapshots"
    headers = {
        "APCA-API-KEY-ID": key,
        "APCA-API-SECRET-KEY": secret
    }

    try:
        r = requests.get(url, headers=headers, params={"symbols": ",".join(symbols)}, timeout=30)

        if r.status_code != 200:
            return jsonify({
                "success": False,
                "message": f"Alpaca market scan failed: {r.status_code} {r.text[:300]}"
            }), 500

        data = r.json()
        snapshots = data.get("snapshots", data)

        rows = []

        for sym in symbols:
            snap = snapshots.get(sym, {}) or {}
            trade = snap.get("latestTrade") or {}
            day = snap.get("dailyBar") or {}

            price = trade.get("p") or day.get("c")
            high = day.get("h")
            low = day.get("l")
            open_price = day.get("o")

            change_pct = None
            signal = "Neutral"

            if price and open_price:
                change_pct = round(((price - open_price) / open_price) * 100, 2)
                if change_pct >= 1:
                    signal = "Bullish"
                elif change_pct <= -1:
                    signal = "Bearish"

            rows.append({
                "symbol": sym,
                "price": price,
                "high": high,
                "low": low,
                "open": open_price,
                "change_pct": change_pct,
                "signal": signal
            })

        rows.append({
            "symbol": "BTC/USD",
            "price": None,
            "high": None,
            "low": None,
            "change_pct": None,
            "signal": "Crypto research mode"
        })

        rows.append({
            "symbol": "GBPJPY",
            "price": None,
            "high": None,
            "low": None,
            "change_pct": None,
            "signal": "OpenAlice forex research mode"
        })

        top = sorted(
            [x for x in rows if isinstance(x.get("change_pct"), (int, float))],
            key=lambda x: x["change_pct"],
            reverse=True
        )[:3]

        return jsonify({
            "success": True,
            "mode": "paper",
            "risk_rule": "1% max stop loss",
            "watchlist": rows,
            "top_3": top,
            "summary": "Here is the move. Capital Leverage scanned the board. Focus on the strongest names, stay in paper mode, and keep risk capped at 1%."
        })

    except Exception as e:
        return jsonify({
            "success": False,
            "message": str(e)
        }), 500


@app.route("/trading/smart-command", methods=["POST"])
def trading_smart_command():
    data = request.get_json(silent=True) or {}
    command = data.get("command", "daily")
    symbol = data.get("symbol", "SPY")

    try:
        scan = trading_market_scan().get_json()
    except Exception as e:
        scan = {"watchlist": [], "top_3": [], "message": str(e)}

    prompt = f"""
You are Capital Leverage Trading Desk.

Selected symbol: {symbol}
Command: {command}

Live market scan:
{scan}

Rules:
- Paper trading only.
- No profit guarantees.
- Max stop loss is 1% per trade.
- Use the live scan to choose the best opportunity.
- Give market-aware guidance.

Return:
HERE IS THE MOVE
BEST OPPORTUNITY RIGHT NOW
TOP 3 TO WATCH
NO-TRADE / AVOID LIST
ENTRY TRIGGER
STOP LOSS PLAN
TAKE PROFIT PLAN
WHAT MARKET/NEWS INFO MATTERS
NEXT ACTION
"""

    try:
        return jsonify(ai_router(prompt, "research"))
    except Exception as e:
        return jsonify({"answer": f"Smart trading command failed: {e}", "provider": "System"}), 500

if __name__ == "__main__":
    app.run(host="0.0.0.0", port=9000, debug=True)
